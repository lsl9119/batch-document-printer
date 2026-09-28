"""
生成 tests/samples/ 下的最小测试文档（内容为本项目自行生成，无第三方版权问题）

用法: python scripts/make_samples.py
依赖: python-docx, openpyxl, python-pptx, Pillow
"""
from pathlib import Path

from PIL import Image, ImageDraw

OUT = Path(__file__).resolve().parent.parent / "tests" / "samples"


def make_docx(path: Path) -> None:
    from docx import Document
    from docx.enum.text import WD_BREAK
    doc = Document()
    doc.add_heading("Batch Document Printer – WPS Edition 测试文档", level=1)
    doc.add_paragraph("第 1 页：中文内容测试。English content test.")
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph("第 2 页：用于验证页数统计与打印。")
    doc.save(path)


def make_xlsx(path: Path) -> None:
    from openpyxl import Workbook
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "A4纵向"
    ws1.page_setup.orientation = "portrait"
    ws1.page_setup.paperSize = ws1.PAPERSIZE_A4
    for row in range(1, 121):  # 多页
        ws1.append([f"行{row}", row, row * 2, "测试数据"])

    ws2 = wb.create_sheet("A3横向")
    ws2.page_setup.orientation = "landscape"
    ws2.page_setup.paperSize = ws2.PAPERSIZE_A3
    for row in range(1, 21):
        ws2.append([f"C{c}-{row}" for c in range(1, 16)])

    ws3 = wb.create_sheet("有打印区域")
    for row in range(1, 51):
        ws3.append([row, f"只打印前10行", row])
    ws3.print_area = "A1:C10"

    ws4 = wb.create_sheet("隐藏表")
    ws4["A1"] = "不应打印"
    ws4.sheet_state = "hidden"
    wb.save(path)


def make_pptx(path: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches
    prs = Presentation()
    for i in range(1, 4):
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = f"幻灯片 {i}"
        slide.placeholders[1].text = "WPS 演示打印测试"
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.save(path)


def _image(text: str, size=(800, 600), color=(30, 120, 200)) -> Image.Image:
    img = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(img)
    draw.rectangle([20, 20, size[0] - 20, size[1] - 20], outline=color, width=8)
    draw.text((60, 60), text, fill=color)
    return img


def make_images() -> None:
    _image("sample.png").save(OUT / "sample.png")
    _image("sample.jpg", color=(200, 60, 30)).save(OUT / "sample.jpg", quality=80)
    _image("sample.bmp", size=(160, 100)).save(OUT / "sample.bmp")
    frames = [_image(f"tiff page {i}", size=(400, 300)) for i in (1, 2)]
    frames[0].save(OUT / "sample.tiff", save_all=True, append_images=frames[1:], compression="tiff_deflate")


def make_pdf(path: Path) -> None:
    pages = [_image(f"PDF page {i}", size=(595, 842)) for i in (1, 2)]
    pages[0].save(path, save_all=True, append_images=pages[1:], resolution=72)


def make_txt() -> None:
    lines = [f"第{i}行：WPS 版批量打印文本测试 English line {i}" for i in range(1, 81)]
    (OUT / "sample.txt").write_text("\n".join(lines), encoding="utf-8")
    (OUT / "sample_gbk.txt").write_bytes("GBK 编码文本：中文测试\n第二行".encode("gbk"))


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    make_docx(OUT / "sample.docx")
    make_xlsx(OUT / "sample.xlsx")
    make_pptx(OUT / "sample.pptx")
    make_pdf(OUT / "sample.pdf")
    make_images()
    make_txt()
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:20} {f.stat().st_size:>8} bytes")


if __name__ == "__main__":
    main()
